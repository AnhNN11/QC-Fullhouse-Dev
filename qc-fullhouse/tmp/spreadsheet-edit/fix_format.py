from pathlib import Path

from openpyxl import load_workbook


ROOT = Path("/Users/nhatanh/qc-fullhouse")
TEMPLATE_PATH = ROOT / "tmp/spreadsheet-edit/Phieu_danh_gia_QC.xlsx"
DATA_PATH = ROOT / "outputs/01a0c74a-0153-7363-b0bc-f42f439112b4/Phieu_danh_gia_QC_26-09-2026.xlsx"
OUTPUT_PATH = ROOT / "outputs/01a0c74a-0153-7363-b0bc-f42f439112b4/Phieu_danh_gia_QC_26-09-2026_format-fixed.xlsx"

CLASS_REVIEWS = [
    {
        "sheet_name": "C++1-1-Thanh Triết",
        "teacher": "Thanh Triết",
        "course": "C++ 1:1",
        "session_no": 3,
        "topic": "Cấu trúc rẽ nhánh",
        "has_title": True,
        "comments": {
            2: "Mic nghe rõ.",
            5: "Nội dung đi đúng phần cấu trúc rẽ nhánh.",
            6: "Giải thích chắc phần điều kiện và cách viết code.",
        },
        "overall": "Chuyên môn tốt, giải thích rõ và có hỏi lại học viên.",
    },
    {
        "sheet_name": "CTDLGT11-Thanh Triết",
        "teacher": "Thanh Triết",
        "course": "CTDLGT 11",
        "session_no": 7,
        "topic": "Ôn tập lý thuyết số, Lũy thừa nhị phân",
        "has_title": True,
        "comments": {
            2: "Mic nghe rõ.",
            5: "Có phân tích lý thuyết số và lũy thừa nhị phân trước khi code.",
            6: "Kiến thức thuật toán chắc, phần nâng cao đi hơi nhanh.",
        },
        "overall": "Chuyên môn tốt. Phần nâng cao đi hơi nhanh.",
    },
    {
        "sheet_name": "Python1-1-Thanh Triết",
        "teacher": "Thanh Triết",
        "course": "Python 1:1",
        "session_no": 26,
        "topic": "Thực hành",
        "has_title": True,
        "comments": {
            2: "Mic nghe rõ.",
            5: "Buổi thực hành bám phần SQL/FastAPI.",
            6: "Giải thích được SQL/FastAPI, xử lý lỗi và cách viết code.",
        },
        "overall": "Hướng dẫn thực hành rõ, có kiểm tra lại mức độ hiểu bài.",
    },
    {
        "sheet_name": "Python1-1-Minh Đạo",
        "teacher": "Minh Đạo",
        "course": "Python 1:1",
        "session_no": 28,
        "topic": None,
        "has_title": False,
        "comments": {},
    },
    {
        "sheet_name": "C++1-1-Vũ Minh Duy",
        "teacher": "Vũ Minh Duy",
        "course": "C++ 1:1",
        "session_no": 22,
        "topic": "Stack và Queue Buổi 1",
        "has_title": True,
        "comments": {},
    },
    {
        "sheet_name": "Python51-Phạm Kiên",
        "teacher": "Phạm Kiên",
        "course": "Python 51",
        "session_no": 13,
        "topic": None,
        "has_title": False,
        "comments": {},
    },
    {
        "sheet_name": "CTDLGT10-Nguyễn Văn Mạnh",
        "teacher": "Nguyễn Văn Mạnh",
        "course": "CTDLGT 10",
        "session_no": 16,
        "topic": "Sinh ke tiep(tt)",
        "has_title": True,
        "comments": {},
    },
    {
        "sheet_name": "BEJava10-Nguyễn Đình Mạnh",
        "teacher": "Nguyễn Đình Mạnh",
        "course": "BE Java 10",
        "session_no": 8,
        "topic": "Abstract Classes & Interfaces",
        "has_title": True,
        "comments": {
            2: "Mic nghe rõ.",
            5: "Nội dung bám phần Abstract Classes & Interfaces.",
            6: "Kiến thức OOP ổn, live-code có một số lỗi cần sửa tại lớp.",
        },
        "overall": "Kiến thức OOP ổn. Cần chuẩn bị code demo kỹ hơn.",
    },
    {
        "sheet_name": "Java33-Nguyễn Đình Mạnh",
        "teacher": "Nguyễn Đình Mạnh",
        "course": "Java 33",
        "session_no": 34,
        "topic": "Projetct 1",
        "has_title": True,
        "comments": {
            2: "Mic nghe rõ.",
            5: "Buổi học đi theo phần đồ án Java Swing/JDBC/MySQL.",
            6: "Giải thích được khung dự án, nhưng phần code và CSDL chuẩn bị chưa kỹ.",
        },
        "overall": "Khung đồ án rõ. Cần chuẩn bị code, CSDL và thiết bị kỹ hơn.",
    },
]

ROW_BY_CRITERION = {
    1: 6, 2: 7, 3: 8, 4: 10, 5: 11, 6: 12, 7: 14, 8: 15, 9: 16,
    10: 18, 11: 19, 12: 21, 13: 22, 14: 23, 15: 24, 16: 26, 17: 27,
}

COMMON_COMMENTS = {
    1: "Không xác định được camera từ file ghi âm.",
    3: "Record không thể hiện rõ giờ bắt đầu và kết thúc.",
    10: "Không kiểm tra được phần chat từ file ghi âm.",
    12: "Buổi này không có phần kiểm tra định kỳ.",
    13: "Buổi này không có phần kiểm tra định kỳ.",
    14: "Buổi này không có phần kiểm tra định kỳ.",
    15: "Buổi này không có phần kiểm tra định kỳ.",
}

NATURAL_COMMENTS = {
    "Thanh Triết": {
        2: "Mic rõ, giữa buổi bị gián đoạn một lần.",
        4: "Dạy đúng nội dung của các lớp C++, CTDLGT và Python.",
        5: "Có phân tích bài trước khi vào code.",
        6: "Kiến thức chắc, giải thích rõ phần thuật toán và SQL/FastAPI.",
        7: "Thường xuyên hỏi lại để kiểm tra học viên có theo kịp không.",
        8: "Có gọi học viên trả lời và kiểm tra code trực tiếp.",
        9: "Xử lý lớp ổn, nhưng có một lần gián đoạn vì việc riêng.",
        11: "Học viên có trả lời bằng mic thường xuyên.",
        16: "Lớp học tích cực, phần nâng cao đi hơi nhanh.",
        17: "Học viên có làm bài. Nên dừng lại kiểm tra thêm ở phần khó.",
    },
    "Minh Đạo": {
        2: "Mic rõ, buổi học không bị ngắt quãng.",
        4: "Dạy đúng phần vòng lặp, range, enumerate, zip và function.",
        5: "Có ôn bài cũ, cho ví dụ rồi để học viên tự làm.",
        6: "Giải thích đúng và khá kỹ các phần dễ nhầm.",
        7: "Câu hỏi gợi mở tốt, không đưa đáp án ngay.",
        8: "Có gọi học viên trả lời xuyên suốt buổi.",
        9: "Kiên nhẫn khi học viên chưa hiểu.",
        11: "Học viên trả lời bằng mic khá thường xuyên.",
        12: "Có nhắc đến bài kiểm tra, nhưng buổi này chưa kiểm tra.",
        16: "Không khí thoải mái, nhưng phần trò chuyện ngoài bài hơi dài.",
        17: "Học viên có phản hồi và theo bài.",
    },
    "Vũ Minh Duy": {
        1: "Camera bị lag.",
        2: "Mạng không ổn định, mic lúc to lúc nhỏ; giữa buổi phải đổi mạng.",
        4: "Dạy đúng phần Stack, Queue, Deque và LIFO/FIFO.",
        5: "Nội dung bám theo mục tiêu học và phần thực hành C++.",
        6: "Giải thích đúng khái niệm và cách dùng.",
        7: "Có hỏi học viên dự đoán kết quả trước khi chạy code.",
        8: "Có trao đổi thường xuyên và điều chỉnh theo tình hình học viên.",
        9: "Kiên nhẫn sửa phần học viên còn nhầm.",
        11: "Học viên có trao đổi bằng mic trong buổi.",
        16: "Không khí thoải mái, ví dụ dễ hiểu.",
        17: "Học viên có tham gia, đầu buổi còn nhầm LIFO/FIFO.",
    },
    "Phạm Kiên": {
        1: "Giáo viên báo camera hỏng nên không bật.",
        2: "Mạng bị rớt giữa buổi.",
        4: "Dạy đúng phần đọc ghi file, đường dẫn và encoding.",
        5: "Có giải thích cách dùng, lỗi thường gặp và cho thực hành.",
        6: "Nắm chắc phần đọc ghi file và các lỗi dễ mất dữ liệu.",
        7: "Có hỏi bài và giao bài tập ngắn.",
        8: "Có gọi học viên và chữa bài gửi qua chat.",
        9: "Vẫn duy trì được lớp khi gặp lỗi camera và mạng.",
        10: "Học viên có gửi đáp án qua chat.",
        11: "Có trao đổi bằng mic với học viên.",
        16: "Lớp học khá tích cực, nhưng bị gián đoạn vì lỗi kỹ thuật.",
        17: "Học viên có làm bài và gửi đáp án.",
    },
    "Nguyễn Văn Mạnh": {
        2: "Mic rõ, nghe được đầy đủ nội dung buổi học.",
        4: "Dạy đúng phần sinh nhị phân, tập con, hoán vị và chỉnh hợp lặp.",
        5: "Có chữa bài cũ, phân tích cách tối ưu rồi cho thực hành.",
        6: "Chuyên môn tốt, giải thích rõ phần TLE và độ phức tạp.",
        7: "Có đặt câu hỏi và yêu cầu học viên giải thích bài.",
        8: "Có gọi nhiều học viên, nhưng cách thúc hơi gắt.",
        9: "Có lời nói khá nặng khi học viên làm sai hoặc nghỉ học. Cần góp ý trực tiếp.",
        11: "Nhiều học viên có trả lời bằng mic.",
        16: "Không khí lớp khá căng ở một số đoạn.",
        17: "Học viên vẫn trả lời, nhưng cách nhắc có thể làm các em áp lực.",
    },
    "Nguyễn Đình Mạnh": {
        2: "Mic rõ, nghe được nội dung của cả hai buổi.",
        4: "Dạy đúng phần OOP và đồ án Java Swing/JDBC/MySQL.",
        5: "Có bám giáo trình và dựng khung cho đồ án.",
        6: "Kiến thức ổn, nhưng lúc live-code có một số lỗi cú pháp và logic.",
        7: "Có hỏi lại để kiểm tra học viên hiểu bài.",
        8: "Có gọi học viên phát biểu và trả lời thắc mắc.",
        9: "Có xử lý được lỗi, nhưng phần code, CSDL và thiết bị chuẩn bị chưa kỹ.",
        11: "Học viên có trao đổi và đặt câu hỏi bằng mic.",
        16: "Lớp trao đổi khá tốt, nhưng lỗi live-code làm chậm nhịp học.",
        17: "Học viên có tham gia thảo luận phần logic đăng nhập.",
    },
}

NATURAL_OVERALL = {
    "Thanh Triết": "Chuyên môn tốt, giải thích kỹ. Phần nâng cao đi hơi nhanh; giữa buổi có một lần gián đoạn.",
    "Minh Đạo": "Giảng kỹ và kiên nhẫn. Nên rút ngắn phần trò chuyện ngoài nội dung bài học.",
    "Vũ Minh Duy": "Tương tác với học viên tốt, ví dụ dễ hiểu. Cần kiểm tra lại camera, mic và đường truyền.",
    "Phạm Kiên": "Bài dạy rõ, có bài tập và chữa bài. Cần sửa camera và kiểm tra đường truyền trước giờ học.",
    "Nguyễn Văn Mạnh": "Chuyên môn tốt. Cần góp ý trực tiếp về cách nói khi học viên làm sai hoặc nghỉ học.",
    "Nguyễn Đình Mạnh": "Ví dụ thực tế và khung đồ án ổn. Cần chuẩn bị code demo, CSDL và thiết bị kỹ hơn.",
}

SESSION_NAME_COMMENTS = {
    "Thanh Triết": (
        "C++ 1:1 - Buổi 3: Cấu trúc rẽ nhánh. "
        "CTDLGT 11 - Buổi 7: Ôn tập lý thuyết số, Lũy thừa nhị phân. "
        "Python 1:1 - Buổi 26: Thực hành."
    ),
    "Minh Đạo": "Python 1:1 - Buổi 28: giáo viên chưa đặt tên buổi học.",
    "Vũ Minh Duy": "C++ 1:1 - Buổi 22: Stack và Queue Buổi 1.",
    "Phạm Kiên": "Python 51 - Buổi 13: giáo viên chưa đặt tên buổi học.",
    "Nguyễn Văn Mạnh": 'CTDLGT 10 - Buổi 16: "Sinh ke tiep(tt)".',
    "Nguyễn Đình Mạnh": (
        'BE Java 10 - Buổi 8: "Abstract Classes & Interfaces". '
        'Java 33 - Buổi 34: "Projetct 1".'
    ),
}

TEACHERS_WITHOUT_SESSION_NAME = {"Minh Đạo", "Phạm Kiên"}
TEACHERS_WITHOUT_CAMERA = {
    "Thanh Triết",
    "Minh Đạo",
    "Phạm Kiên",
    "Nguyễn Văn Mạnh",
    "Nguyễn Đình Mạnh",
}


template_book = load_workbook(TEMPLATE_PATH)
data_book = load_workbook(DATA_PATH, data_only=False)
template_sheet = template_book["Mau danh gia QC"]

for review in CLASS_REVIEWS:
    teacher = review["teacher"]
    source = data_book[teacher]
    target = template_book.copy_worksheet(template_sheet)
    target.title = review["sheet_name"]

    # Chỉ ghi nội dung mới vào bản sao của sheet mẫu. Toàn bộ style, merge,
    # border, fill, font, alignment, kích thước hàng/cột và thiết lập in được
    # giữ nguyên từ worksheet gốc.
    target["A1"] = f"PHIẾU ĐÁNH GIÁ QC - {review['course'].upper()} - {teacher.upper()}"
    topic_label = review["topic"] if review["has_title"] else "Chưa đặt tên buổi học"
    target["A2"] = f"Ngày 26/09/2026 - {review['course']} - Buổi {review['session_no']}: {topic_label}"

    for criterion, row in ROW_BY_CRITERION.items():
        for column in range(3, 8):
            if criterion == 1 and teacher in TEACHERS_WITHOUT_CAMERA:
                target.cell(row=row, column=column).value = "X" if column == 3 else None
            elif criterion == 4 and not review["has_title"]:
                target.cell(row=row, column=column).value = None
            else:
                target.cell(row=row, column=column).value = source.cell(row=row, column=column).value
        comment = (
            "Giáo viên không bật camera."
            if criterion == 1 and teacher in TEACHERS_WITHOUT_CAMERA
            else f"{review['course']} - Buổi {review['session_no']}: {review['topic']}."
            if criterion == 4
            and review["has_title"]
            else f"{review['course']} - Buổi {review['session_no']}: giáo viên chưa đặt tên buổi học."
            if criterion == 4
            else review["comments"][criterion]
            if criterion in review["comments"]
            else NATURAL_COMMENTS[teacher].get(criterion, COMMON_COMMENTS.get(criterion))
        )
        target.cell(row=row, column=8).value = comment

    for address in ("A33", "E33", "A35", "E35"):
        target[address] = source[address].value

    target["A30"] = review.get("overall", NATURAL_OVERALL[teacher])

    # Không ghi thông tin công cụ phân tích vào phiếu QC gửi nội bộ.
    target["A31"] = None

    target.sheet_view.showGridLines = template_sheet.sheet_view.showGridLines

template_book.active = template_book.sheetnames.index("Mau danh gia QC")
OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
template_book.save(OUTPUT_PATH)

print(OUTPUT_PATH)
print(template_book.sheetnames)
